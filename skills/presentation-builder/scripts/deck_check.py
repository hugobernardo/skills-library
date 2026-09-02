#!/usr/bin/env python3
"""
Mechanical pre-flight for a .pptx, to run before the visual QA pass.

This catches defects that are unambiguous in geometry: shapes off-canvas, text
bleeding past the slide boundary, text blocks genuinely overlapping, content
crowding a slide edge, and slides missing presenter notes. It also inventories
unfilled placeholders so none reach the room by accident.

Two things it deliberately does NOT do, because getting them wrong is worse
than not doing them:

  - Text overflow. Estimating rendered text extent without a layout engine
    produces confident wrong answers, and overflow is the most common real
    defect. It stays the job of the visual pass.
  - Proximity between every pair of shapes. A label sits close to the number it
    labels on purpose. Flagging that as crowding buries the real findings and
    teaches you to ignore the tool.

So: run this to clear the mechanical faults, then render and look at every slide.

    python deck_check.py deck.pptx [--edge 0.5]

Exit code is 0 when nothing is wrong, 1 when an error is found. Warnings and
placeholders do not fail the check — warnings need your eyes on the render, and
placeholders are usually deliberate.
"""
import sys, re, argparse
from pptx import Presentation
from pptx.util import Emu

PLACEHOLDER = re.compile(
    r"\[\s*(FILL|TBD|TK|XX+|INSERT|PLACEHOLDER|SOURCE|DATA|POSTMORTEM|NAME|LOGO)[^\]]*\]|lorem ipsum", re.I)


def inches(v):
    return round(Emu(int(v)).inches, 3) if v is not None else None


def walk(shapes):
    """Every shape, recursing into groups. Text inside a group is invisible to a
    naive top-level iteration, which makes slides look empty when they are not."""
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:  # GROUP
            yield from walk(sh.shapes)


def check(path, edge_min=0.5):
    prs = Presentation(path)
    W, H = inches(prs.slide_width), inches(prs.slide_height)
    errors, warnings, placeholders = [], [], []

    for i, slide in enumerate(prs.slides, 1):
        text_boxes = []
        for sh in walk(slide.shapes):
            if sh.left is None or sh.top is None:
                continue
            l, t = inches(sh.left), inches(sh.top)
            w, h = inches(sh.width) or 0, inches(sh.height) or 0
            r, b = l + w, t + h
            has_text = sh.has_text_frame and sh.text_frame.text.strip()
            label = (sh.text_frame.text.strip().split("\n")[0][:40] if has_text else str(sh.shape_type))

            if r < 0 or b < 0 or l > W or t > H:
                errors.append((i, "off-canvas", f"'{label}' sits entirely outside the slide"))
                continue

            if has_text:
                if l < -0.02 or t < -0.02 or r > W + 0.02 or b > H + 0.02:
                    errors.append((i, "past edge", f"'{label}' extends past the slide boundary"))
                else:
                    near = min(l, t, W - r, H - b)
                    if near < edge_min - 0.02:
                        warnings.append((i, "tight margin",
                                         f"'{label}' is {near:.2f}in from a slide edge (want {edge_min}in)"))
                    text_boxes.append((l, t, r, b, label))

                m = PLACEHOLDER.search(sh.text_frame.text)
                if m:
                    placeholders.append((i, m.group(0).strip()))

        # Possible overlaps. Reported as warnings, not errors: a text box is
        # routinely taller than the text inside it, so bounding boxes intersect
        # on slides that render perfectly. Require the intersection to be a real
        # share of the smaller box before saying anything at all.
        for a in range(len(text_boxes)):
            for c in range(a + 1, len(text_boxes)):
                l1, t1, r1, b1, n1 = text_boxes[a]
                l2, t2, r2, b2, n2 = text_boxes[c]
                ox = min(r1, r2) - max(l1, l2)
                oy = min(b1, b2) - max(t1, t2)
                if ox <= 0.1 or oy <= 0.25:
                    continue
                if oy < 0.45 * min(b1 - t1, b2 - t2):
                    continue
                warnings.append((i, "possible overlap",
                                 f"'{n1}' and '{n2}' intersect by {ox:.2f} x {oy:.2f}in — confirm in the render"))

        if not (slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip()):
            errors.append((i, "no notes", "slide has no presenter notes"))

    return errors, warnings, placeholders, len(prs.slides)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx")
    ap.add_argument("--edge", type=float, default=0.5, help="minimum margin from a slide edge, inches")
    a = ap.parse_args()

    errors, warnings, placeholders, n = check(a.pptx, a.edge)

    print(f"{a.pptx} — {n} slides")

    def block(title, items):
        if not items:
            return
        print(f"\n{title}")
        last = None
        for slide, kind, msg in sorted(items):
            if slide != last:
                print(f"  Slide {slide}")
                last = slide
            print(f"    [{kind}] {msg}")

    block(f"ERRORS ({len(errors)})", errors)
    block(f"WARNINGS ({len(warnings)})", warnings)

    if placeholders:
        print(f"\nUNFILLED PLACEHOLDERS ({len(placeholders)}) — deliberate if you put them there, "
              f"but none of these can reach the room:")
        for slide, marker in sorted(placeholders):
            print(f"  Slide {slide}: {marker}")

    if not errors and not warnings:
        print("\nNo mechanical faults found.")
    print("\nNow render and look at every slide. Text overflow, contrast and visual "
          "balance are not detectable here.")
    return 1 if errors else 0


if __name__ == "__main__":
    sys.exit(main())
